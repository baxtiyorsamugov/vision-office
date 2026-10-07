"""Read-only attendance reports built from locally recorded camera events."""

from __future__ import annotations

from collections import defaultdict
from dataclasses import dataclass
from datetime import date, datetime, timedelta
from io import BytesIO
from pathlib import Path
from xml.sax.saxutils import escape

from sqlalchemy.orm import Session

from core.local_time import as_utc, device_timezone, local_day_bounds_utc, to_local
from database.models import RecognitionEvent


REPORT_TYPES = ("eduschool_employee", "eduschool_student")
MAX_DAYS = 31
MAX_EVENTS = 250_000
MAX_VISIT_HOURS = 18


@dataclass
class ReportEvent:
    id: str
    person_id: str
    person_type: str
    name: str
    camera_id: str
    direction: str
    occurred_at: datetime
    local_at: datetime
    mark: str = ""


@dataclass(frozen=True)
class ReportVisit:
    person_id: str
    person_type: str
    name: str
    entry_at: datetime | None
    exit_at: datetime | None
    status: str
    day: date

    @property
    def minutes(self) -> int | None:
        if self.entry_at is None or self.exit_at is None:
            return None
        return round((as_utc(self.exit_at) - as_utc(self.entry_at)).total_seconds() / 60)


@dataclass(frozen=True)
class PersonReportRow:
    person_id: str
    person_type: str
    name: str
    days: int
    visits: int
    completed_minutes: int
    first_entry: datetime | None
    last_exit: datetime | None
    incomplete: int


@dataclass(frozen=True)
class DailyReportRow:
    day: date
    employees: int
    students: int
    completed_visits: int
    observations: int


@dataclass
class AttendanceReport:
    start: date
    end: date
    generated_at: datetime
    events: list[ReportEvent]
    visits: list[ReportVisit]
    people: list[PersonReportRow]
    days: list[DailyReportRow]

    @property
    def employee_count(self) -> int:
        return sum(row.person_type != "eduschool_student" for row in self.people)

    @property
    def student_count(self) -> int:
        return sum(row.person_type == "eduschool_student" for row in self.people)

    @property
    def completed_count(self) -> int:
        return sum(visit.status == "complete" for visit in self.visits)

    @property
    def incomplete_count(self) -> int:
        return sum(visit.status != "complete" for visit in self.visits)


def _source(person_type: str) -> str:
    return "Локальный" if person_type == "local_employee" else "EduSchool"


def _direction(value: str) -> str:
    return "Вход" if value == "entry" else "Выход"


def _local_text(value: datetime | None, pattern: str = "%d.%m %H:%M") -> str:
    return to_local(value).strftime(pattern) if value else "—"


def _duration(minutes: int | None) -> str:
    return f"{minutes // 60:02d}:{minutes % 60:02d}" if minutes is not None else "—"


def _pair_person_events(events: list[ReportEvent]) -> list[ReportVisit]:
    visits: list[ReportVisit] = []
    opened: ReportEvent | None = None
    last_exit: ReportEvent | None = None
    for event in events:
        if event.direction == "entry":
            if opened is not None:
                gap = as_utc(event.occurred_at) - as_utc(opened.occurred_at)
                if gap <= timedelta(hours=MAX_VISIT_HOURS) and event.local_at.date() == opened.local_at.date():
                    event.mark = "Повторная фиксация"
                    continue
                opened.mark = "Вход без выхода"
                visits.append(ReportVisit(opened.person_id, opened.person_type, opened.name,
                                          opened.occurred_at, None, "no_exit", opened.local_at.date()))
            opened = event
            last_exit = None
            event.mark = "Вход без выхода"
            continue

        if opened is not None:
            elapsed = as_utc(event.occurred_at) - as_utc(opened.occurred_at)
            if timedelta(0) <= elapsed <= timedelta(hours=MAX_VISIT_HOURS):
                opened.mark = "Вход визита"
                event.mark = "Выход визита"
                visits.append(ReportVisit(opened.person_id, opened.person_type, opened.name,
                                          opened.occurred_at, event.occurred_at, "complete", opened.local_at.date()))
                opened = None
                last_exit = event
                continue
            opened.mark = "Вход без выхода"
            visits.append(ReportVisit(opened.person_id, opened.person_type, opened.name,
                                      opened.occurred_at, None, "no_exit", opened.local_at.date()))
            opened = None

        if last_exit is not None and as_utc(event.occurred_at) - as_utc(last_exit.occurred_at) <= timedelta(minutes=5):
            event.mark = "Повторная фиксация"
        else:
            event.mark = "Выход без входа"
            visits.append(ReportVisit(event.person_id, event.person_type, event.name,
                                      None, event.occurred_at, "no_entry", event.local_at.date()))
            last_exit = event
    if opened is not None:
        opened.mark = "Вход без выхода"
        visits.append(ReportVisit(opened.person_id, opened.person_type, opened.name,
                                  opened.occurred_at, None, "no_exit", opened.local_at.date()))
    return visits


def load_attendance_report(engine, start: date, end: date, *, include_local: bool = False,
                           max_events: int = MAX_EVENTS) -> AttendanceReport:
    """Fetch only report columns; never load photos, embeddings, or confidence."""
    if start > end or (end - start).days >= MAX_DAYS:
        raise ValueError(f"Выберите период от 1 до {MAX_DAYS} дней.")
    start_utc, _ = local_day_bounds_utc(start)
    _, end_utc = local_day_bounds_utc(end)
    kinds = REPORT_TYPES + (("local_employee",) if include_local else ())
    timezone = device_timezone()
    groups: dict[tuple[str, str], list[ReportEvent]] = defaultdict(list)
    with Session(engine) as session:
        rows = session.query(
            RecognitionEvent.id, RecognitionEvent.person_id, RecognitionEvent.person_type,
            RecognitionEvent.person_name, RecognitionEvent.camera_id,
            RecognitionEvent.event_type, RecognitionEvent.created_at,
        ).filter(
            RecognitionEvent.created_at >= start_utc - timedelta(hours=MAX_VISIT_HOURS),
            RecognitionEvent.created_at < end_utc + timedelta(hours=MAX_VISIT_HOURS),
            RecognitionEvent.person_type.in_(kinds),
            RecognitionEvent.event_type.in_(("entry", "exit")),
        ).order_by(RecognitionEvent.person_type, RecognitionEvent.person_id,
                   RecognitionEvent.created_at, RecognitionEvent.id).yield_per(1000)
        for index, row in enumerate(rows):
            if index >= max_events:
                raise ValueError("За период слишком много фиксаций. Сократите диапазон дат.")
            person_id = row.person_id or f"missing:{row.id}"
            groups[(row.person_type, person_id)].append(ReportEvent(
                id=row.id, person_id=person_id, person_type=row.person_type,
                name=row.person_name or "Без имени", camera_id=row.camera_id,
                direction=row.event_type, occurred_at=as_utc(row.created_at),
                local_at=as_utc(row.created_at).astimezone(timezone),
            ))

    all_events = [event for group in groups.values() for event in group]
    all_visits = [visit for group in groups.values() for visit in _pair_person_events(group)]
    events = sorted((event for event in all_events if start <= event.local_at.date() <= end),
                    key=lambda item: (item.occurred_at, item.id))
    visits = sorted((visit for visit in all_visits if start <= visit.day <= end),
                    key=lambda item: (item.entry_at or item.exit_at, item.person_id))

    visits_by_person: dict[tuple[str, str], list[ReportVisit]] = defaultdict(list)
    for visit in visits:
        visits_by_person[(visit.person_type, visit.person_id)].append(visit)
    events_by_person: dict[tuple[str, str], list[ReportEvent]] = defaultdict(list)
    for event in events:
        events_by_person[(event.person_type, event.person_id)].append(event)
    people = []
    for key in set(events_by_person) | set(visits_by_person):
        person_events = events_by_person[key]
        person_visits = visits_by_person[key]
        entries = [event.occurred_at for event in person_events if event.direction == "entry"]
        exits = [event.occurred_at for event in person_events if event.direction == "exit"]
        people.append(PersonReportRow(
            person_id=key[1], person_type=key[0],
            name=(person_events[-1].name if person_events else person_visits[-1].name),
            days=len({event.local_at.date() for event in person_events}),
            visits=sum(visit.status == "complete" for visit in person_visits),
            completed_minutes=sum(visit.minutes or 0 for visit in person_visits if visit.status == "complete"),
            first_entry=min(entries) if entries else None,
            last_exit=max(exits) if exits else None,
            incomplete=sum(visit.status != "complete" for visit in person_visits),
        ))
    people.sort(key=lambda person: (person.person_type == "eduschool_student", person.name.casefold(), person.person_id))

    daily_events: dict[date, list[ReportEvent]] = defaultdict(list)
    for event in events:
        daily_events[event.local_at.date()].append(event)
    daily_complete: dict[date, int] = defaultdict(int)
    for visit in visits:
        if visit.status == "complete":
            daily_complete[visit.day] += 1
    days = []
    for offset in range((end - start).days + 1):
        day = start + timedelta(days=offset)
        observed = daily_events[day]
        days.append(DailyReportRow(
            day=day,
            employees=len({(event.person_type, event.person_id) for event in observed
                           if event.person_type != "eduschool_student"}),
            students=len({event.person_id for event in observed if event.person_type == "eduschool_student"}),
            completed_visits=daily_complete[day], observations=len(observed),
        ))
    return AttendanceReport(start, end, datetime.now().astimezone(), events, visits, people, days)


def _pdf_fonts() -> tuple[str, str]:
    from reportlab.pdfbase import pdfmetrics
    from reportlab.pdfbase.ttfonts import TTFont

    locations = (
        (Path("/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf"),
         Path("/usr/share/fonts/truetype/dejavu/DejaVuSans-Bold.ttf")),
        (Path("C:/Windows/Fonts/arial.ttf"), Path("C:/Windows/Fonts/arialbd.ttf")),
    )
    for regular, bold in locations:
        if regular.is_file() and bold.is_file():
            if "ReportSans" not in pdfmetrics.getRegisteredFontNames():
                pdfmetrics.registerFont(TTFont("ReportSans", str(regular)))
                pdfmetrics.registerFont(TTFont("ReportSans-Bold", str(bold)))
            return "ReportSans", "ReportSans-Bold"
    raise RuntimeError("Для PDF нужен шрифт с кириллицей в Docker или Windows Fonts.")


def build_pdf(report: AttendanceReport, *, branch_name: str = "Филиал EduSchool") -> bytes:
    from reportlab.lib import colors
    from reportlab.lib.enums import TA_CENTER
    from reportlab.lib.pagesizes import A4, landscape
    from reportlab.lib.styles import ParagraphStyle
    from reportlab.lib.units import mm
    from reportlab.platypus import LongTable, Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

    font, bold = _pdf_fonts()
    green = colors.HexColor("#108455")
    ink = colors.HexColor("#172923")
    muted = colors.HexColor("#68776f")
    paper = colors.HexColor("#f4f7f5")
    output = BytesIO()
    page_size = landscape(A4)
    doc = SimpleDocTemplate(output, pagesize=page_size, leftMargin=18*mm, rightMargin=18*mm,
                            topMargin=22*mm, bottomMargin=16*mm, title="Отчёт по посещениям")
    title = ParagraphStyle("report-title", fontName=bold, fontSize=19, leading=24, textColor=ink)
    subtitle = ParagraphStyle("report-subtitle", fontName=font, fontSize=9, leading=13, textColor=muted)
    section = ParagraphStyle("report-section", fontName=bold, fontSize=12, leading=16, textColor=ink,
                             spaceBefore=16, spaceAfter=7)
    cell = ParagraphStyle("report-cell", fontName=font, fontSize=7.6, leading=10, textColor=ink)
    cell_bold = ParagraphStyle("report-cell-bold", parent=cell, fontName=bold)
    cell_header = ParagraphStyle("report-cell-header", parent=cell_bold, textColor=colors.white)
    center = ParagraphStyle("report-center", parent=cell, alignment=TA_CENTER)
    width = page_size[0] - doc.leftMargin - doc.rightMargin

    def paragraph(value, style=cell):
        return Paragraph(escape(str(value)), style)

    def table(rows, widths, header=True):
        result = LongTable(rows, colWidths=widths, repeatRows=1 if header else 0, hAlign="LEFT")
        commands = [
            ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
            ("LEFTPADDING", (0, 0), (-1, -1), 7), ("RIGHTPADDING", (0, 0), (-1, -1), 7),
            ("TOPPADDING", (0, 0), (-1, -1), 5), ("BOTTOMPADDING", (0, 0), (-1, -1), 5),
            ("LINEBELOW", (0, -1), (-1, -1), 0.5, colors.HexColor("#dce8e1")),
        ]
        if header:
            commands += [("BACKGROUND", (0, 0), (-1, 0), green),
                         ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
                         ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, paper])]
        result.setStyle(TableStyle(commands))
        return result

    story = [Paragraph("Отчёт по посещениям", title), Spacer(1, 3*mm),
             Paragraph(escape(branch_name[:100]), subtitle),
             Paragraph(f"Период: {report.start:%d.%m.%Y} - {report.end:%d.%m.%Y} · "
                       f"Сформирован: {_local_text(report.generated_at, '%d.%m.%Y %H:%M')} · Asia/Tashkent", subtitle),
             Spacer(1, 8*mm)]
    metrics = [
        ("Сотрудники", report.employee_count), ("Ученики", report.student_count),
        ("Завершённые визиты", report.completed_count), ("Неполные данные", report.incomplete_count),
    ]
    metric_cells = [Table([[paragraph(label, subtitle)], [Paragraph(str(value), ParagraphStyle(
        f"metric-{index}", fontName=bold, fontSize=17, leading=21, textColor=green))]],
        colWidths=[width/4 - 12]) for index, (label, value) in enumerate(metrics)]
    story.append(table([metric_cells], [width/4]*4, header=False))
    story.append(Paragraph("По дням", section))
    day_rows = [[paragraph(value, cell_header) for value in
                 ("Дата", "Сотрудники", "Ученики", "Завершённые визиты", "Фиксации")]]
    day_rows.extend([[paragraph(day.day.strftime("%d.%m.%Y")), paragraph(day.employees, center),
                      paragraph(day.students, center), paragraph(day.completed_visits, center),
                      paragraph(day.observations, center)] for day in report.days])
    story.append(table(day_rows, [width*.20]*5))

    for label, is_student in (("Сотрудники", False), ("Ученики", True)):
        people = [person for person in report.people if (person.person_type == "eduschool_student") == is_student]
        story.append(Paragraph(f"{label} · {len(people)}", section))
        headers = ("№", "ФИО", "Источник", "Дней", "Визитов", "Время, ч:м", "Первый вход", "Последний выход", "Неполные")
        rows = [[paragraph(value, cell_header) for value in headers]]
        for index, person in enumerate(people, 1):
            values = (index, person.name, _source(person.person_type), person.days, person.visits,
                      _duration(person.completed_minutes) if person.visits else "—",
                      _local_text(person.first_entry), _local_text(person.last_exit), person.incomplete)
            rows.append([paragraph(value) for value in values])
        if not people:
            rows.append([paragraph("—"), paragraph("За период нет фиксаций")] + [paragraph("")]*7)
        story.append(table(rows, [width*.04, width*.24, width*.09, width*.06, width*.07,
                                  width*.10, width*.15, width*.15, width*.10]))
    story.append(Spacer(1, 7*mm))
    story.append(Paragraph("Источник: локальные события камер. «Не зафиксирован» не означает «отсутствовал». "
                           "Время пребывания считается только для пар вход–выход. Полный журнал фиксаций — в Excel.", subtitle))

    def decorate(canvas, document):
        canvas.saveState()
        canvas.setStrokeColor(colors.HexColor("#dce8e1"))
        canvas.line(doc.leftMargin, page_size[1] - 14*mm, page_size[0] - doc.rightMargin, page_size[1] - 14*mm)
        canvas.setFont(font, 8)
        canvas.setFillColor(muted)
        canvas.drawString(doc.leftMargin, 9*mm, "Vision Office · локальный отчёт")
        canvas.drawRightString(page_size[0] - doc.rightMargin, 9*mm, f"Стр. {document.page}")
        canvas.restoreState()

    doc.build(story, onFirstPage=decorate, onLaterPages=decorate)
    return output.getvalue()


def _excel_text(value: str) -> str:
    """Keep external names/IDs as text, never as spreadsheet formulas."""
    return "'" + value if value and value[0] in "=+-@" else value


def build_excel(report: AttendanceReport, *, branch_name: str = "Филиал EduSchool") -> bytes:
    from openpyxl import Workbook
    from openpyxl.cell import WriteOnlyCell
    from openpyxl.styles import Alignment, Font, PatternFill

    workbook = Workbook(write_only=True)
    green = PatternFill("solid", fgColor="108455")
    white = Font(name="Calibri", size=10, bold=True, color="FFFFFF")
    normal = Font(name="Calibri", size=10, color="172923")
    date_format = "dd.mm.yyyy hh:mm:ss"
    timezone = device_timezone()

    def sheet(title: str, headers: tuple[str, ...], widths: tuple[int, ...]):
        ws = workbook.create_sheet(title)
        for index, width in enumerate(widths, 1):
            ws.column_dimensions[chr(64 + index)].width = width
        ws.freeze_panes = "A2"
        header_cells = []
        for heading in headers:
            cell = WriteOnlyCell(ws, value=heading)
            cell.fill = green
            cell.font = white
            cell.alignment = Alignment(vertical="center")
            header_cells.append(cell)
        ws.append(header_cells)
        return ws

    def append(ws, values):
        cells = []
        for value in values:
            if isinstance(value, str):
                value = _excel_text(value)
            cell = WriteOnlyCell(ws, value=value)
            cell.font = normal
            if isinstance(value, datetime):
                cell.number_format = date_format
            cells.append(cell)
        ws.append(cells)

    def local_datetime(value):
        return as_utc(value).astimezone(timezone).replace(tzinfo=None) if value else None

    overview = sheet("Сводка", ("Показатель", "Значение"), (35, 28))
    for label, value in (
        ("Филиал", branch_name[:100]), ("Начало", report.start.strftime("%d.%m.%Y")),
        ("Окончание", report.end.strftime("%d.%m.%Y")),
        ("Сформирован", _local_text(report.generated_at, "%d.%m.%Y %H:%M")),
        ("Сотрудники", report.employee_count), ("Ученики", report.student_count),
        ("Завершённые визиты", report.completed_count), ("Неполные данные", report.incomplete_count),
        ("Фиксации", len(report.events)),
    ):
        append(overview, (label, value))
    daily = sheet("По дням", ("Дата", "Сотрудники", "Ученики", "Завершённые визиты", "Фиксации"),
                  (17, 18, 18, 25, 18))
    for row in report.days:
        append(daily, (row.day.strftime("%d.%m.%Y"), row.employees, row.students,
                       row.completed_visits, row.observations))

    people_headers = ("ФИО", "Источник", "ID", "Дней с фиксацией", "Завершённые визиты",
                      "Время, ч:м", "Первый вход", "Последний выход", "Неполные данные")
    employee_sheet = sheet("Сотрудники", people_headers, (40, 18, 38, 21, 22, 16, 24, 24, 20))
    student_sheet = sheet("Ученики", people_headers, (40, 18, 38, 21, 22, 16, 24, 24, 20))
    for person in report.people:
        append(student_sheet if person.person_type == "eduschool_student" else employee_sheet,
               (person.name, _source(person.person_type), person.person_id, person.days, person.visits,
                _duration(person.completed_minutes) if person.visits else None,
                local_datetime(person.first_entry), local_datetime(person.last_exit), person.incomplete))

    visits_sheet = sheet("Визиты", ("ФИО", "Тип", "Источник", "ID", "Дата визита", "Вход", "Выход",
                                    "Время, ч:м", "Состояние"), (40, 18, 18, 38, 17, 24, 24, 16, 25))
    visit_status = {"complete": "Завершён", "no_exit": "Нет выхода", "no_entry": "Нет входа"}
    for visit in report.visits:
        append(visits_sheet, (visit.name, "Ученик" if visit.person_type == "eduschool_student" else "Сотрудник",
                              _source(visit.person_type), visit.person_id, visit.day.strftime("%d.%m.%Y"),
                              local_datetime(visit.entry_at), local_datetime(visit.exit_at),
                              _duration(visit.minutes), visit_status[visit.status]))

    event_sheet = sheet("Все фиксации", ("Время", "ФИО", "Тип", "Источник", "ID человека", "Направление",
                                       "Камера", "Отметка", "ID события"),
                        (24, 40, 18, 18, 38, 18, 24, 25, 40))
    for event in report.events:
        append(event_sheet, (event.local_at.replace(tzinfo=None), event.name,
                             "Ученик" if event.person_type == "eduschool_student" else "Сотрудник",
                             _source(event.person_type), event.person_id, _direction(event.direction),
                             event.camera_id, event.mark, event.id))
    output = BytesIO()
    workbook.save(output)
    return output.getvalue()
