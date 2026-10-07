import tempfile
import unittest
from datetime import date, datetime, timedelta, timezone
from io import BytesIO
from pathlib import Path

from sqlalchemy import create_engine
from sqlalchemy.orm import Session

from core.attendance_reports import build_excel, build_pdf, load_attendance_report
from database.migrations import run_migrations
from database.models import RecognitionEvent


class AttendanceReportTests(unittest.TestCase):
    def setUp(self):
        self.tempdir = tempfile.TemporaryDirectory()
        self.engine = create_engine(f"sqlite:///{Path(self.tempdir.name) / 'reports.db'}")
        run_migrations(self.engine)

    def tearDown(self):
        self.engine.dispose()
        self.tempdir.cleanup()

    def event(self, number, person_type, person_id, direction, hour, minute=0, name="Same Name"):
        with Session(self.engine) as session:
            session.add(RecognitionEvent(
                id=f"00000000-0000-0000-0000-{number:012d}",
                camera_id="entry" if direction == "entry" else "exit",
                event_type=direction, person_id=person_id, person_type=person_type,
                person_name=name, subject_signature=person_id,
                created_at=datetime(2026, 10, 6, hour, minute, tzinfo=timezone.utc),
                confidence=0.93,
            ))
            session.commit()

    def test_every_recorded_event_is_retained_but_repeat_is_not_a_visit(self):
        self.event(1, "eduschool_employee", "edu:e:abc", "entry", 3)
        self.event(2, "eduschool_employee", "edu:e:abc", "entry", 3, 2)
        self.event(3, "eduschool_employee", "edu:e:abc", "exit", 8)
        self.event(4, "eduschool_employee", "edu:e:abc", "exit", 8, 1)
        self.event(5, "eduschool_employee", "edu:e:abc", "entry", 9)
        self.event(6, "eduschool_student", "edu:s:xyz", "entry", 4)
        self.event(7, "eduschool_student", "edu:s:xyz", "exit", 5)
        self.event(8, "eduschool_employee", "edu:e:other", "exit", 6)
        self.event(9, "local_employee", "local:7", "entry", 3)
        day = date(2026, 10, 6)
        report = load_attendance_report(self.engine, day, day)
        self.assertEqual(len(report.events), 8)
        self.assertEqual(report.employee_count, 2)
        self.assertEqual(report.student_count, 1)
        self.assertEqual(report.completed_count, 2)
        self.assertEqual(report.incomplete_count, 2)
        self.assertEqual(report.days[0].observations, 8)
        marks = {event.id[-1]: event.mark for event in report.events}
        self.assertEqual(marks["2"], "Повторная фиксация")
        self.assertEqual(marks["4"], "Повторная фиксация")
        self.assertEqual(marks["5"], "Вход без выхода")
        self.assertEqual(marks["8"], "Выход без входа")
        staff = next(person for person in report.people if person.person_id == "edu:e:abc")
        self.assertEqual(staff.visits, 1)
        self.assertEqual(staff.completed_minutes, 300)
        self.assertEqual(staff.incomplete, 1)
        self.assertEqual(staff.days, 1)
        with_local = load_attendance_report(self.engine, day, day, include_local=True)
        self.assertEqual(len(with_local.events), 9)
        self.assertEqual(with_local.employee_count, 3)

    def test_overnight_visit_is_assigned_to_entry_day(self):
        self.event(1, "eduschool_student", "edu:s:night", "entry", 18, 30)
        self.event(2, "eduschool_student", "edu:s:night", "exit", 20)
        first = load_attendance_report(self.engine, date(2026, 10, 6), date(2026, 10, 6))
        self.assertEqual(first.completed_count, 1)
        self.assertEqual(first.visits[0].minutes, 90)
        self.assertEqual(len(first.events), 1)
        self.assertEqual(first.events[0].mark, "Вход визита")
        second = load_attendance_report(self.engine, date(2026, 10, 7), date(2026, 10, 7))
        self.assertEqual(second.completed_count, 0)
        self.assertEqual(len(second.events), 1)
        self.assertEqual(second.events[0].mark, "Выход визита")

    def test_export_has_all_events_without_confidence_or_formulas(self):
        self.event(1, "eduschool_employee", "edu:e:abc", "entry", 3, name="=TEST()")
        self.event(2, "eduschool_employee", "edu:e:abc", "exit", 8, name="=TEST()")
        report = load_attendance_report(self.engine, date(2026, 10, 6), date(2026, 10, 6))
        self.assertTrue(build_pdf(report, branch_name="Филиал Один").startswith(b"%PDF"))
        from openpyxl import load_workbook
        book = load_workbook(BytesIO(build_excel(report)), read_only=True)
        self.assertEqual(book.sheetnames, ["Сводка", "По дням", "Сотрудники", "Ученики", "Визиты", "Все фиксации"])
        events = list(book["Все фиксации"].values)
        self.assertEqual(len(events), 3)
        self.assertEqual(events[0][0], "Время")
        self.assertNotIn("confidence", str(events[0]).lower())
        self.assertNotIn("Уверенность", str(events[0]))
        self.assertEqual(events[1][1], "'=TEST()")
        self.assertEqual(book["Сводка"]["B6"].value, 1)
        self.assertEqual(book["Сотрудники"]["F2"].value, "05:00")
        book.close()

    def test_empty_report_and_limits(self):
        report = load_attendance_report(self.engine, date(2026, 10, 6), date(2026, 10, 6))
        self.assertEqual(report.completed_count, 0)
        self.assertTrue(build_pdf(report).startswith(b"%PDF"))
        self.assertTrue(build_excel(report).startswith(b"PK"))
        with self.assertRaises(ValueError):
            load_attendance_report(self.engine, date(2026, 10, 7), date(2026, 10, 6))
        with self.assertRaises(ValueError):
            load_attendance_report(self.engine, date(2026, 9, 1), date(2026, 10, 7))
        self.event(1, "eduschool_student", "edu:s:abc", "entry", 4)
        with self.assertRaises(ValueError):
            load_attendance_report(self.engine, date(2026, 10, 6), date(2026, 10, 6), max_events=0)


if __name__ == "__main__":
    unittest.main()
